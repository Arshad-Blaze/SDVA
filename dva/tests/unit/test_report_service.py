import polars as pl
import pytest

from models.report_models import ItemValidationConfig
from models.validation_models import ColumnMapping, PriceType, UnitType
from services.reporting.report_service import ReportError, generate_reports


def write_parquet(tmp_path, name, data: pl.DataFrame):
    path = tmp_path / name
    data.write_parquet(path)
    return path


def item_config(**kwargs):
    kwargs.setdefault(
        "bau",
        ColumnMapping(upc_col="UPC", desc_col="DESC", units_col="UNITS", price_col="PRICE"),
    )
    kwargs.setdefault(
        "test",
        ColumnMapping(upc_col="UPC", desc_col="DESC", units_col="UNITS", price_col="PRICE"),
    )
    return ItemValidationConfig(**kwargs)


def make_pair(tmp_path):
    """BAU: 3 rows over 2 UPCs; TEST: 2 rows, one shared + one new UPC."""
    bau = write_parquet(
        tmp_path,
        "bau.parquet",
        pl.DataFrame(
            {
                "UPC": ["11111", "11111", "22222"],
                "DESC": ["aaa", "aaa", "bbb"],
                "UNITS": [2, 3, 5],
                "PRICE": [10, 10, 20],
            }
        ),
    )
    test = write_parquet(
        tmp_path,
        "test.parquet",
        pl.DataFrame(
            {
                "UPC": ["11111", "33333"],
                "DESC": ["aaa", "ccc"],
                "UNITS": [1, 7],
                "PRICE": [10, 5],
            }
        ),
    )
    return bau, test


def run(tmp_path, config=None):
    bau, test = make_pair(tmp_path)
    out = tmp_path / "reports"
    result = generate_reports(bau, test, config or item_config(), out)
    return result, out


def test_bau_summary_aggregates_by_upc_desc(tmp_path):
    result, out = run(tmp_path)
    bau = pl.read_parquet(out / "item_validation_bau_summary.parquet")
    assert bau.height == 2
    assert set(bau["UPC"].to_list()) == {"11111", "22222"}
    row = bau.filter(pl.col("UPC") == "11111")
    assert row["Product_Description"][0] == "aaa"
    assert row["TOTAL_UNITS"][0] == 5.0
    # default TOTAL_PRICE: dollars = sum of the price column (10 + 10)
    assert row["TOTAL_DOLLARS"][0] == 20.0


def test_comparison_classifies_presence(tmp_path):
    result, out = run(tmp_path)
    comp = pl.read_parquet(out / "item_validation_comparison.parquet")
    by_upc = {row["UPC"]: row["Present In"] for row in comp.iter_rows(named=True)}
    assert by_upc["11111"] == "Present in Both"
    assert by_upc["22222"] == "Present only in BAU"
    assert by_upc["33333"] == "Present only in TEST"


def test_comparison_differences(tmp_path):
    result, out = run(tmp_path)
    comp = pl.read_parquet(out / "item_validation_comparison.parquet")
    shared = comp.filter(pl.col("UPC") == "11111")
    # BAU units 5 vs TEST units 1 -> +4; dollars 20 vs 10 -> +10
    assert shared["Units Difference"][0] == 4.0
    assert shared["Dollar Difference"][0] == 10.0


def test_summary_groups_and_grand_total(tmp_path):
    result, out = run(tmp_path)
    summary = pl.read_parquet(out / "item_validation_summary.parquet")
    assert sorted(summary["Present In"].to_list())[:3] == [
        "Grand Total",
        "Present in Both",
        "Present only in BAU",
    ]
    assert "Present only in TEST" in summary["Present In"].to_list()
    assert summary.filter(pl.col("Present In") == "Present in Both")["Count"][0] == 1
    grand = summary.filter(pl.col("Present In") == "Grand Total")
    assert grand["Count"][0] == 3


def test_metrics_and_top_bottom_frames(tmp_path):
    result, out = run(tmp_path)
    metrics = {row["Metric"]: row["Value"] for row in result.metrics.iter_rows(named=True)}
    assert metrics["Unique UPC BAU"] == 2
    assert metrics["Unique UPC TEST"] == 2
    assert metrics["New UPC indicator"] == 0
    assert metrics["Total BAU Sales"] == 40.0  # 20 + 20 (TOTAL_PRICE)
    assert metrics["Total Test Sales"] == 15.0  # 10 + 5
    assert metrics["Units Difference"] == 2.0  # BAU 10 vs TEST 8
    assert result.top_5_sales.height == 3  # fewer than 5 rows is fine
    assert result.top_5_units.height == 3


def test_unit_price_and_implied_decimals(tmp_path):
    bau, test = make_pair(tmp_path)
    config = item_config(price_type_bau=PriceType.UNIT_PRICE, implied_dollars_bau=True)
    out = tmp_path / "reports"
    result = generate_reports(bau, test, config, out)
    bau = pl.read_parquet(out / "item_validation_bau_summary.parquet")
    row = bau.filter(pl.col("UPC") == "11111")
    # dollars = units * price / 100 = (2*10 + 3*10)/100 = 0.5
    assert row["TOTAL_DOLLARS"][0] == 0.5


def test_weighted_units_override(tmp_path):
    bau = write_parquet(
        tmp_path,
        "bau.parquet",
        pl.DataFrame(
            {
                "UPC": ["11111", "11111"],
                "DESC": ["aaa", "aaa"],
                "UNITS": [2, 3],
                "WTD": [None, 10],
                "PRICE": [10, 10],
            }
        ),
    )
    test = write_parquet(
        tmp_path,
        "test.parquet",
        pl.DataFrame(
            {"UPC": ["11111"], "DESC": ["aaa"], "UNITS": [1], "WTD": [2], "PRICE": [10]}
        ),
    )
    config = ItemValidationConfig(
        bau=ColumnMapping(upc_col="UPC", desc_col="DESC", units_col="UNITS", price_col="PRICE", weighted_units_col="WTD"),
        test=ColumnMapping(upc_col="UPC", desc_col="DESC", units_col="UNITS", price_col="PRICE", weighted_units_col="WTD"),
    )
    out = tmp_path / "reports"
    generate_reports(bau, test, config, out)
    bau_summary = pl.read_parquet(out / "item_validation_bau_summary.parquet")
    # weighted 10 replaces the 3 (second row), first row keeps units 2 -> total 12
    assert bau_summary["TOTAL_UNITS"][0] == 12.0


def test_weighted_units_drive_unit_price_dollars(tmp_path):
    bau = write_parquet(
        tmp_path,
        "bau.parquet",
        pl.DataFrame(
            {
                "UPC": ["11111", "11111"],
                "DESC": ["aaa", "aaa"],
                "UNITS": [2, 3],
                "WTD": [None, 10],
                "PRICE": [10, 10],
            }
        ),
    )
    test = write_parquet(
        tmp_path,
        "test.parquet",
        pl.DataFrame(
            {"UPC": ["11111"], "DESC": ["aaa"], "UNITS": [1], "WTD": [2], "PRICE": [10]}
        ),
    )
    config = ItemValidationConfig(
        bau=ColumnMapping(upc_col="UPC", desc_col="DESC", units_col="UNITS", price_col="PRICE", weighted_units_col="WTD"),
        test=ColumnMapping(upc_col="UPC", desc_col="DESC", units_col="UNITS", price_col="PRICE", weighted_units_col="WTD"),
        price_type_bau=PriceType.UNIT_PRICE,
    )
    out = tmp_path / "reports"
    generate_reports(bau, test, config, out)
    bau_summary = pl.read_parquet(out / "item_validation_bau_summary.parquet")
    # dollars = price x effective units: 10*2 + 10*10 = 120
    assert bau_summary["TOTAL_DOLLARS"][0] == 120.0
    assert bau_summary["TOTAL_UNITS"][0] == 12.0


def test_unit_type_column_in_outputs(tmp_path):
    bau, test = make_pair(tmp_path)
    config = item_config(
        bau=ColumnMapping(upc_col="UPC", desc_col="DESC", units_col="UNITS", price_col="PRICE", units_type=UnitType.WEIGHT),
    )
    out = tmp_path / "reports"
    generate_reports(bau, test, config, out)
    assert pl.read_parquet(out / "item_validation_bau_summary.parquet")["Unit Type"].to_list() == ["weight", "weight"]
    comparison = pl.read_parquet(out / "item_validation_comparison.parquet")
    assert set(comparison["Unit Type"].to_list()) == {"weight"}
    summary = pl.read_parquet(out / "item_validation_summary.parquet")
    assert set(summary["Unit Type"].to_list()) == {"weight"}


def test_incomplete_config_raises(tmp_path):
    bau, test = make_pair(tmp_path)
    config = item_config(
        bau=ColumnMapping(upc_col="UPC", desc_col="DESC", units_col="UNITS"),
    )
    with pytest.raises(ValueError, match="incomplete"):
        generate_reports(bau, test, config, tmp_path / "reports")


def test_missing_dataset_raises(tmp_path):
    with pytest.raises(ReportError, match="not found"):
        generate_reports(
            tmp_path / "nope.parquet",
            tmp_path / "also.parquet",
            item_config(),
            tmp_path / "reports",
        )


def test_artifacts_are_readable_parquet(tmp_path):
    result, out = run(tmp_path)
    assert set(result.artifacts) == {
        "bau_summary",
        "test_summary",
        "comparison",
        "summary",
        "excel",
    }
    for artifact in result.artifacts.values():
        assert artifact.exists()
        if artifact.suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(artifact, read_only=True)
            assert workbook.sheetnames == ["BAU Summary", "TEST Summary", "Comparison", "Summary"]
            # the comparison sheet has its header plus all comparison rows
            headers = next(workbook["Comparison"].iter_rows(values_only=True))
            assert "Present In" in headers
            continue
        df = pl.read_parquet(artifact)
        assert df.height >= 1
    assert result.elapsed_seconds >= 0.0
